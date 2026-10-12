from pathlib import Path
import csv,json,hashlib,urllib.request,subprocess,io,zipfile,sys
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
W=Path(__file__).resolve().parent;R=W.parents[2];T=R/'Reference/tables/96_金融下載與分頁分類續核_B74_20261012';T.mkdir(parents=True,exist_ok=True)
def read(p):return list(csv.DictReader(p.open(encoding='utf-8-sig')))
def write(name,rows):
 fields=list(dict.fromkeys(k for x in rows for k in x)) or ['無記錄']
 with (T/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def get(item):
 u=item['target_url'];a=dict(網址=u,來源觀察數=len(item['observations']),核查日='2026-10-12',來源觀察=json.dumps(item['observations'],ensure_ascii=False),完整性='UNKNOWN（未知）');rows=[];members=[]
 try:
  with urllib.request.urlopen(urllib.request.Request(u,headers={'User-Agent':'ResearchDirectoryContent/1.0'}),timeout=20) as f:
   b=f.read(20000001);a.update(HTTP狀態=f.status,最終網址=f.url,內容型別=f.headers.get('Content-Type',''),位元組=len(b),截斷=len(b)>20000000,SHA256=hashlib.sha256(b).hexdigest())
  p=W/'原始回應'/ (hashlib.sha256(u.encode()).hexdigest()[:20]+'.response');p.parent.mkdir(exist_ok=True);p.write_bytes(b);a['原始檔']=p.relative_to(R).as_posix()
  low=u.lower();labels=' '.join(x.get('label','') for x in item['observations']).lower()
  role='其他資料待核'
  if 'historical trade' in labels:role='個別證券歷史交易，非公司全量'
  elif 'full list of securities' in labels:role='證券清單候選，須分類公司／基金／債券／衍生品'
  elif 'brokers facilitating' in labels:role='網上開戶券商子集，非全體券商'
  elif 'isin' in labels:role='證券識別碼清單，非獨立公司數'
  elif 'round lot' in labels:role='交易手數規則清單，非公司全量'
  elif 'form' in labels or 'questionnaire' in labels or 'checklist' in labels or 'compliance' in labels:role='申請或合規表單，非機構名冊'
  elif 'company_annoucements' in low:role='公司公告分頁，非公司名冊'
  elif 'instruments-rules-policies' in low:role='法規政策分頁，非公司名冊'
  a['業務分類']=role
  if b.startswith(b'PK'):
   z=zipfile.ZipFile(io.BytesIO(b));members=[dict(網址=u,成員名稱=x.filename,壓縮位元組=x.compress_size,原始位元組=x.file_size) for x in z.infolist()]
   if '[Content_Types].xml' in z.namelist() and 'xl/workbook.xml' in z.namelist():
    book=load_workbook(io.BytesIO(b),read_only=True,data_only=False)
    for s in book:
     for n,values in enumerate(s.iter_rows(values_only=True),1):
      if any(v is not None for v in values):rows.append(dict(網址=u,工作表=s.title,行=n,欄值JSON=json.dumps(values,ensure_ascii=False,default=str),語義='來源行，尚未核定為實體'))
    book.close();a['解析']='XLSX所有非空來源行；公式保留，不執行'
   else:a['解析']='ZIP僅核成員目錄，不執行或自動解壓'
  elif '.csv' in low:
   for n,v in enumerate(csv.reader(io.StringIO(b.decode('utf-8-sig'))),1):rows.append(dict(網址=u,工作表='CSV',行=n,欄值JSON=json.dumps(v,ensure_ascii=False),語義='來源行，尚未核定為實體'))
   a['解析']='CSV全部來源行'
  else:a['解析']='原頁或舊式XLS保留，後續解析待核'
 except Exception as e:a['錯誤類型']=type(e).__name__
 return a,rows,members
assert not (W/'結果.json').exists()
B=R/'Reference/tables/95_全球金融未嘗試來源續查_B73_20261012';paths=[B/'下載連結觀察.csv',B/'分頁連結觀察.csv'];base={p.relative_to(R).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
(W/'基線.json').write_text(json.dumps({'HEAD':subprocess.check_output(['git','rev-parse','HEAD'],cwd=R).decode().strip(),'來源':base,'Git狀態':subprocess.check_output(['git','status','--short'],cwd=R).decode('utf-8')},ensure_ascii=False,indent=2),encoding='utf-8')
obs=sum([read(p) for p in paths],[]);group={}
for x in obs:group.setdefault(x['target_url'],[]).append(x)
items=[dict(target_url=u,observations=v) for u,v in group.items()];results=[];rows=[];members=[]
with ThreadPoolExecutor(max_workers=4) as pool:
 for a,b,c in pool.map(get,items):results.append(a);rows+=b;members+=c;print('已核',len(results),flush=True)
write('下載與分頁收據.csv',results);write('試算表來源全量行.csv',rows);write('壓縮檔成員清單.csv',members)
assert len(results)==len(group)
for p in paths:assert hashlib.sha256(p.read_bytes()).hexdigest()==base[p.relative_to(R).as_posix()]
for a in results:
 if a.get('原始檔'):assert hashlib.sha256((R/a['原始檔']).read_bytes()).hexdigest()==a['SHA256']
summary=dict(原觀察=len(obs),唯一網址=len(results),HTTP200=sum(a.get('HTTP狀態')==200 for a in results),來源行=len(rows),壓縮成員=len(members),完整性升級=0)
(W/'結果.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
text='---\ntitle: "金融下載與分頁分類續核：第七十四批"\nlang: zh-Hant\nformat: html\n---\n\n保留第七十三批全部下載及分頁觀察，按完整網址去重後逐一取得。來源行包括標題、表頭及不同證券類型，不是公司或券商數。業務分類依連結標籤與來源用途初審，逐實體關聯仍待核實。\n\n'+json.dumps(summary,ensure_ascii=False,indent=2)+'\n\n[下載與分頁收據](tables/96_金融下載與分頁分類續核_B74_20261012/下載與分頁收據.csv)；[試算表全量來源行](tables/96_金融下載與分頁分類續核_B74_20261012/試算表來源全量行.csv)。\n\nZIP只列成員目錄；XLSX公式保留不執行；舊式XLS、公告及法規原頁後續另核。截斷、錯誤及授權未核狀態逐項保留。全球完整性沒有升級，後續先對證券清單做實體分類、日期核對及去重，再核券商子集與壓縮檔內容。\n'
(R/'Reference/金融下載與分頁分類續核_B74_20261012.qmd').write_text(text,encoding='utf-8');print(json.dumps(summary,ensure_ascii=False))
