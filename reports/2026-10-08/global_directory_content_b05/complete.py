import intake,csv,json,hashlib,sqlite3,io,zipfile,urllib.request
from lxml import html
from concurrent.futures import ThreadPoolExecutor
O,T,R=intake.O,intake.T,intake.R
def read(n):
    return list(csv.DictReader((T/n).open(encoding='utf-8-sig')))
manifest=read('directory_fetch_manifest.csv')
cma=next(r for r in manifest if 'cmauganda.co.ug/cma-licensed-firms' in r['url'])
tree=html.fromstring((R/cma['file']).read_bytes());cards=[];sections=[]
for section in tree.xpath('//div[contains(concat(" ",normalize-space(@class)," ")," cma-section ")]'):
    category=' '.join(section.xpath('.//h2')[0].text_content().split())
    nodes=section.xpath('.//div[contains(concat(" ",normalize-space(@class)," ")," cma-card ")]')
    counts=section.xpath('./div[1]//span/text()');declared=next((int(x.strip()) for x in counts if x.strip().isdigit()),None)
    sections.append(dict(category=category,declared_count=declared,observed_cards=len(nodes),count_reconciled=declared==len(nodes),section_html=html.tostring(section,encoding='unicode')))
    for i,node in enumerate(nodes,1):
        children=node.xpath('./div');name=' '.join(children[0].text_content().split()) if children else ''
        cards.append(dict(source_url=cma['url'],category=category,ordinal=i,name=name,card_text=' '.join(node.text_content().split()),card_html=html.tostring(node,encoding='unicode'),evidence='OFFICIAL_CATEGORY_CARD_CURRENT_RENEWAL_NOT_INDEPENDENTLY_VERIFIED'))
intake.write('registry_cma_all_licensed_category_cards.csv',cards);intake.write('registry_cma_section_count_reconciliation.csv',sections)
urls=sorted({r['target_url'] for r in read('registry_directory_observed_download_links.csv')});D=O/'downloads';D.mkdir(exist_ok=True)
def fetch(url):
    result=dict(url=url,checked_on='2026-10-08');members=[];rows=[]
    try:
        if 'List-of-Online-Brokers-13-8-24.zip' in url:
            data=(R/'reports/2026-10-08/global_market_directories_b04/downloaded/file_04.response').read_bytes();result['acquisition']='REUSE_PRESERVED_B04_BYTES'
        else:
            with urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'ResearchDirectory/1.0'}),timeout=30) as response:
                data=response.read(20000001);result['http_status']=response.status
            result['acquisition']='DIRECT_OFFICIAL_LINK'
        result.update(bytes=len(data),sha256=hashlib.sha256(data).hexdigest(),truncated=len(data)>20000000)
        path=D/(hashlib.sha256(url.encode()).hexdigest()[:20]+'.response');path.write_bytes(data);result['file']=path.relative_to(R).as_posix()
        if zipfile.is_zipfile(io.BytesIO(data)):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for info in z.infolist():
                    if info.is_dir():continue
                    if info.file_size>50000000:
                        members.append(dict(url=url,member=info.filename,bytes=info.file_size,status='SIZE_LIMIT_NOT_EXPANDED'));continue
                    content=z.read(info)
                    members.append(dict(url=url,member=info.filename,bytes=len(content),sha256=hashlib.sha256(content).hexdigest(),status='FULL_MEMBER_OBSERVED'))
        if url.lower().endswith('.xlsx'):
            import openpyxl
            workbook=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
            for sheet in workbook:
                for number,values in enumerate(sheet.iter_rows(values_only=True),1):
                    rows.append(dict(url=url,sheet=sheet.title,row=number,cells_json=json.dumps(list(values),ensure_ascii=False,default=str),semantic_status='SOURCE_WORKBOOK_ROW_NOT_AUTOMATICALLY_COMPANY'))
    except Exception as e:result['error']=type(e).__name__
    return result,members,rows
results=[];members=[];rows=[]
with ThreadPoolExecutor(max_workers=5) as pool:
    for result,m,r in pool.map(fetch,urls):results.append(result);members.extend(m);rows.extend(r)
intake.write('direct_file_fetch_manifest.csv',results);intake.write('registry_download_archive_members.csv',members);intake.write('registry_download_workbook_rows.csv',rows)
summary=dict(cma_cards=len(cards),cma_sections=len(sections),cma_section_count_mismatches=sum(not r['count_reconciled'] for r in sections),unique_download_links=len(urls),downloads_saved=sum('file' in r for r in results),archive_members=len(members),workbook_rows=len(rows),global_complete=False)
(O/'completion_acceptance.json').write_text(json.dumps(summary,indent=2)+'\n',encoding='utf-8');print(json.dumps(summary))
