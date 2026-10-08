from pathlib import Path
import csv,json,re,hashlib,io,zipfile
from urllib.request import urlopen,Request
from urllib.error import HTTPError
from openpyxl import load_workbook
O=Path(__file__).parent;R=O.parents[2];T=R/'Reference/tables/29_global_market_directories_b04_20261008';D=O/'downloaded';D.mkdir(exist_ok=True)
def read(name):return list(csv.DictReader((T/name).open(encoding='utf-8-sig')))
observations=read('registry_observed_directory_links.csv')+read('registry_anna_observed_directory_links.csv')
urls=sorted({r['target_url'] for r in observations if re.search(r'[.](csv|xlsx?|zip)([?]|$)',r['target_url'],re.I)})
manifest=[];records=[]
for number,url in enumerate(urls,1):
    result=dict(url=url,observed_from_json=json.dumps([r for r in observations if r['target_url']==url],ensure_ascii=False),retrieved_on='2026-10-08',semantic_status='SOURCE_FILE_NOT_ASSUMED_COMPANY_OR_BROKER_DIRECTORY')
    try:
        with urlopen(Request(url,headers={'User-Agent':'ResearchDirectoryInventory/1.0'}),timeout=12) as response:
            b=response.read(10000001);result.update(http_status=response.status,final_url=response.url,bytes=len(b),sha256=hashlib.sha256(b).hexdigest(),truncated=len(b)>10000000)
        path=D/f'file_{number:02d}.response';path.write_bytes(b);result['file']=path.relative_to(R).as_posix()
        if result['truncated']:manifest.append(result);continue
        if b[:2]==b'PK' and '.zip' in url.lower():
            with zipfile.ZipFile(io.BytesIO(b)) as archive:
                contents=[]
                for name in archive.namelist():
                    info=archive.getinfo(name);content=archive.read(name) if info.file_size<50000000 else None
                    contents.append(dict(member=name,bytes=info.file_size,sha256=hashlib.sha256(content).hexdigest() if content is not None else 'NOT_READ_SIZE_LIMIT'))
                result['archive_members_json']=json.dumps(contents)
        elif b[:2]==b'PK':
            workbook=load_workbook(io.BytesIO(b),read_only=True,data_only=False);count=0
            for sheet in workbook:
                for line,values in enumerate(sheet.iter_rows(values_only=True),1):
                    records.append(dict(url=url,sheet=sheet.title,source_row=line,raw_cells_json=json.dumps(values,ensure_ascii=False,default=str)));count+=1
            result['all_workbook_rows']=count;workbook.close()
        elif '.csv' in url.lower():
            text=b.decode('utf-8-sig',errors='backslashreplace');count=0
            for line,values in enumerate(csv.reader(io.StringIO(text)),1):
                records.append(dict(url=url,source_row=line,raw_cells_json=json.dumps(values,ensure_ascii=False)));count+=1
            result['all_CSV_rows']=count
    except HTTPError as e:result.update(http_status=e.code,error='HTTP_ERROR')
    except Exception as e:result.update(error=type(e).__name__)
    manifest.append(result)
def write(name,values):
    fields=list(dict.fromkeys(k for v in values for k in v))
    with (T/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(values)
write('observed_data_file_manifest.csv',manifest);write('registry_observed_file_all_rows.csv',records)
print(json.dumps([{k:v for k,v in row.items() if k not in ('observed_from_json','archive_members_json')} for row in manifest]),flush=True)
for row in manifest:
    if row.get('archive_members_json'):print(row['archive_members_json'],flush=True)
