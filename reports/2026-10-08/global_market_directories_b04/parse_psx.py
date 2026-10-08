from pathlib import Path
import csv,json,io,zipfile,hashlib
from openpyxl import load_workbook
O=Path(__file__).parent;R=O.parents[2];T=R/'Reference/tables/29_global_market_directories_b04_20261008'
archive=zipfile.ZipFile(io.BytesIO((O/'downloaded/file_04.response').read_bytes()));rows=[]
for member in archive.namelist():
    if not member.lower().endswith('.xlsx'):continue
    b=archive.read(member);book=load_workbook(io.BytesIO(b),read_only=True,data_only=False)
    for sheet in book:
        for number,values in enumerate(sheet.iter_rows(values_only=True),1):
            rows.append(dict(source_member=member,sheet=sheet.title,source_row=number,raw_cells_json=json.dumps(values,ensure_ascii=False,default=str),
                source_date_from_filename='2024-08-13_NOT_CURRENT_LIST_VERIFIED',retrieved_on='2026-10-08',member_sha256=hashlib.sha256(b).hexdigest()))
    book.close()
with (T/'registry_psx_online_broker_all_workbook_rows.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
brokers=[]
for row in rows:
    cells=json.loads(row['raw_cells_json'])
    if row['source_row']==1 or len(cells)<8 or not cells[2]:continue
    brokers.append(dict(source_row=row['source_row'],member_code=cells[1],member_name=cells[2],status_in_2024_source=cells[3],
                        address=cells[4],phone=cells[5],email=cells[6],website=cells[7],country='PK',
                        source_date='2024-08-13_FILENAME_DATE',retrieved_on='2026-10-08',current_authorisation='UNKNOWN',
                        coverage='ONLINE_BROKER_SOURCE_SNAPSHOT_NOT_ALL_PAKISTAN_LICENSEES',raw_cells_json=row['raw_cells_json']))
with (T/'registry_psx_online_broker_snapshot.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(brokers[0]));w.writeheader();w.writerows(brokers)
(O/'psx_parse_acceptance.json').write_text(json.dumps(dict(all_workbook_rows=len(rows),member_source_rows=len(brokers),unique_member_codes=len({r['member_code'] for r in brokers}),current_full_register_verified=False),indent=2)+'\n',encoding='utf-8')
print(json.dumps(dict(all_workbook_rows=len(rows),first_rows=rows[:5]),ensure_ascii=True))
