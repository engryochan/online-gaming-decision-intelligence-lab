from pathlib import Path
import sys,csv,json,hashlib,re,io,zipfile,urllib.request,collections
from urllib.parse import urlparse,quote
from concurrent.futures import ThreadPoolExecutor
O=Path(__file__).resolve().parent;R=O.parents[2];T=R/'Reference/tables/42_global_directory_followup_b17_20261008'
sys.path.insert(0,str(O.parent/'global_directory_content_b05'));import intake
intake.O=O;intake.T=T
assert not (O/'download_acceptance.json').exists()
def read(p):return list(csv.DictReader(p.open(encoding='utf-8-sig')))
links=read(T/'registry_download_links.csv')+read(T/'registry_pagination_download_links.csv');byurl=collections.defaultdict(list)
for r in links:byurl[r['target_url']].append(r)
prior={}
for batch in ['30_global_directory_content_b05_20261008','31_global_directory_content_b06_20261008','32_global_directory_content_b07_20261008','35_global_directory_content_b10_20261008','38_global_directory_followup_b13_20261008','40_global_directory_followup_b15_20261008']:
    for r in read(R/'Reference/tables'/batch/'direct_file_fetch_manifest.csv'):
        if r.get('file'):prior[r['url']]=r
montreal=sorted([u for u in byurl if urlparse(u).hostname=='www.m-x.ca']);other=[u for u in byurl if u not in montreal]
selected=set(byurl)
selected.update(u for u in montreal if 'options-summary' in u)
queue=[dict(target_url=u,observations=len(v),source_observations_json=json.dumps(v,ensure_ascii=False),action='SELECTED_THIS_BATCH' if u in selected else 'PENDING_FULL_HISTORICAL_FILE_ACQUISITION',global_complete='UNKNOWN') for u,v in sorted(byurl.items())]
intake.write('registry_all_download_workqueue.csv',queue)
D=O/'downloads';D.mkdir(exist_ok=True)
def fetch(url):
    result=dict(url=url,checked_on='2026-10-08');members=[];rows=[];csvrows=[]
    try:
        if url in prior:
            r=prior[url];data=(R/r['file']).read_bytes();assert hashlib.sha256(data).hexdigest()==r['sha256']
            result.update(file=r['file'],acquisition='REUSE_PRIOR_HASH_VERIFIED',http_status=r.get('http_status',''),truncated=r.get('truncated')=='True')
        else:
            requested=quote(url,safe=':/?=&%#+');result['requested_url']=requested
            with urllib.request.urlopen(urllib.request.Request(requested,headers={'User-Agent':'ResearchDirectory/1.0'}),timeout=30) as response:
                data=response.read(20000001);result.update(http_status=response.status,final_url=response.url,content_type=response.headers.get('Content-Type',''))
            result.update(acquisition='DIRECT_OBSERVED_LINK',truncated=len(data)>20000000)
            p=D/(hashlib.sha256(url.encode()).hexdigest()[:20]+'.response');p.write_bytes(data);result['file']=p.relative_to(R).as_posix()
        result.update(bytes=len(data),sha256=hashlib.sha256(data).hexdigest())
        if result['truncated']:result['parse_status']='BODY_LIMIT_PENDING';return result,members,rows,csvrows
        if zipfile.is_zipfile(io.BytesIO(data)):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for info in z.infolist():
                    if info.is_dir():continue
                    if info.file_size>50000000:members.append(dict(url=url,member=info.filename,bytes=info.file_size,status='SIZE_LIMIT_PENDING'));continue
                    content=z.read(info);members.append(dict(url=url,member=info.filename,bytes=len(content),sha256=hashlib.sha256(content).hexdigest(),status='MEMBER_OBSERVED_NOT_AUTOMATICALLY_BUSINESS_ENTITY'))
        path=urlparse(url).path.lower()
        if path.endswith('.xlsx'):
            import openpyxl
            wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
            for sheet in wb:
                for n,values in enumerate(sheet.iter_rows(values_only=True),1):rows.append(dict(url=url,sheet=sheet.title,row=n,cells_json=json.dumps(list(values),ensure_ascii=False,default=str),role='SOURCE_WORKBOOK_ROW_SEMANTICS_PENDING'))
            wb.close()
        elif path.endswith('.csv'):
            try:text=data.decode('utf-8-sig');encoding='UTF8'
            except UnicodeDecodeError:text=data.decode('latin1');encoding='LATIN1_REVERSIBLE_ENCODING_UNCONFIRMED'
            sample=text[:8192]
            try:delimiter=csv.Sniffer().sniff(sample,delimiters=',;\t|').delimiter
            except csv.Error:delimiter='; '[:1] if sample.count(';')>sample.count(',') else ','
            for n,cells in enumerate(csv.reader(io.StringIO(text),delimiter=delimiter),1):csvrows.append(dict(source_url=url,row=n,encoding=encoding,delimiter=delimiter,cells_json=json.dumps(cells,ensure_ascii=False),role='SOURCE_CSV_ROW_SECURITY_EVENT_OR_OTHER_SCOPE_REVIEW'))
    except urllib.error.HTTPError as e:result.update(error='HTTPError',http_status=e.code)
    except Exception as e:result['error']=type(e).__name__
    return result,members,rows,csvrows
results=[];members=[];rows=[];csvrows=[]
with ThreadPoolExecutor(max_workers=8) as pool:
    for result,m,r,c in pool.map(fetch,sorted(selected)):results.append(result);members.extend(m);rows.extend(r);csvrows.extend(c)
intake.write('direct_file_fetch_manifest.csv',results);intake.write('registry_download_archive_members.csv',members);intake.write('registry_download_workbook_rows.csv',rows);intake.write('registry_download_csv_rows.csv',csvrows)
summary=dict(download_observations=len(links),unique_links=len(byurl),selected_links=len(selected),saved_or_reused=sum(bool(r.get('file')) for r in results),reused=sum(r.get('acquisition')=='REUSE_PRIOR_HASH_VERIFIED' for r in results),new_saved=sum(r.get('acquisition')=='DIRECT_OBSERVED_LINK' for r in results),pending_unselected=len(byurl)-len(selected),remaining_errors=sum(bool(r.get('error')) for r in results),archive_members=len(members),workbook_rows=len(rows),csv_rows=len(csvrows),global_complete=False)
(O/'download_acceptance.json').write_text(json.dumps(summary,indent=2)+'\n',encoding='utf-8');print(json.dumps(summary))
